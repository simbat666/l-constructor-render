#!/usr/bin/env python3
"""Create a versioned JSON rule snapshot from an asynchronous-motor OL workbook.

The workbook is a build-time engineering source. The web application and DXF
service use the generated JSON, so a deployed instance never needs Excel or an
absolute path on the engineer's computer. Question columns are located by their
headers: adding a column such as ``Required field`` cannot silently shift rules.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "data" / "rules-source" / "ОЛ - Электродвигатель асинхронный_В2 [nsCQfC].xlsx"
DEFAULT_TARGET = ROOT / "data" / "motor-v2.json"


def value(cell):
    if cell is None:
        return None
    if isinstance(cell, float) and cell.is_integer():
        return int(cell)
    return cell


def num(cell):
    return float(cell) if cell is not None else None


def rows(ws):
    # Artifact-tool exports omit the optional XLSX <dimension> element.  Force
    # openpyxl to determine the occupied range so iter_rows() remains bounded.
    if ws.max_row is None or ws.max_column is None:
        ws.calculate_dimension(force=True)
    return list(ws.iter_rows(values_only=True))


def headers(ws) -> dict[str, int]:
    if ws.max_row is None or ws.max_column is None:
        ws.calculate_dimension(force=True)
    result = {str(cell.value).strip(): index for index, cell in enumerate(ws[1]) if cell.value is not None}
    if not result:
        raise ValueError(f"Лист {ws.title} не содержит строку заголовков")
    return result


def question_rows(ws) -> list[dict[str, object]]:
    columns = headers(ws)
    required_headers = {"order", "Вопрос", "binds", "Answer", "Ответ", "Field type"}
    missing = sorted(required_headers - set(columns))
    if missing:
        raise ValueError(f"Лист {ws.title}: отсутствуют обязательные колонки: {', '.join(missing)}")

    required_column = columns.get("Required field")
    engine_key_column = columns.get("Engine key")
    diagram_key_column = columns.get("key_diagram")
    result = []
    for source_row, row in enumerate(rows(ws)[1:], start=2):
        if row[columns["order"]] is None:
            continue
        field_type = value(row[columns["Field type"]])
        input_rule = value(row[columns["Answer"]]) if field_type == "input" else None
        if field_type == "input" and not re.match(r"^(?:float|string)\|", str(input_rule or "")):
            raise ValueError(
                f"Лист {ws.title}, строка {source_row}: для input отсутствует техническое правило float или string"
            )
        result.append({
            "sourceRow": source_row,
            "order": value(row[columns["order"]]),
            "question": value(row[columns["Вопрос"]]),
            "binds": value(row[columns["binds"]]),
            "answer": value(row[columns["Ответ"]]),
            "fieldType": field_type,
            "inputRule": input_rule,
            "engineKey": value(row[engine_key_column]) if engine_key_column is not None else None,
            "diagramKey": value(row[diagram_key_column]) if diagram_key_column is not None else None,
            "required": bool(value(row[required_column])) if required_column is not None else False,
        })
    return result


def validate_question_graph(question_sets: list[list[dict[str, object]]]) -> None:
    """Reject broken or regressed OL dependencies before they enter the product."""
    nodes = [node for question_set in question_sets for node in question_set]
    orders = [str(node["order"]) for node in nodes]
    duplicates = sorted({order for order in orders if orders.count(order) > 1})
    if duplicates:
        raise ValueError("Повторяющиеся коды вопросов: " + ", ".join(duplicates))
    known_orders = set(orders)
    missing = []
    for node in nodes:
        targets = [item.strip() for item in str(node.get("binds") or "").split(";") if item.strip()]
        if len(targets) != len(set(targets)):
            raise ValueError(f"Вопрос {node['order']} содержит повторяющуюся связь")
        for target in targets:
            if target not in known_orders:
                missing.append(f"{node['order']} → {target}")
    if missing:
        raise ValueError("Ссылки ведут на несуществующие вопросы: " + ", ".join(missing))


IO_CODES = ("DI24-HSC-NPN", "DI24-NPN", "DI24-PNP", "DOT-PNP", "DOR-NO",
            "AI-RTD2", "AI-I0_20-PAS", "AI-I4_20-PAS", "AI-U0_10",
            "AO-U0_10", "RS-485")


def required_columns(ws, names):
    columns = headers(ws)
    missing = sorted(set(names) - set(columns))
    if missing:
        raise ValueError(f"Лист {ws.title}: отсутствуют колонки: {', '.join(missing)}")
    return columns


def cell(row, columns, name):
    return value(row[columns[name]]) if name in columns else None


def io(row, columns, source_row):
    result = {}
    for code in IO_CODES:
        raw = cell(row, columns, code)
        if isinstance(raw, bool) or not isinstance(raw, (int, float)) or raw < 0 or int(raw) != raw:
            raise ValueError(f"Лист I/O, строка {source_row}: {code} должен быть целым неотрицательным числом")
        result[code] = int(raw)
    return result


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Extract a production rule snapshot from a motor OL workbook")
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE, help="Source .xlsx workbook")
    parser.add_argument("--target", type=Path, default=DEFAULT_TARGET, help="Target JSON rule snapshot")
    parser.add_argument("--revision", help="Human-readable source revision stored in JSON; defaults to source stem")
    return parser.parse_args()


def main():
    args = parse_args()
    source = args.source.expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(f"Не найден исходный ОЛ: {source}")
    book = load_workbook(source, read_only=True, data_only=True)

    questions1 = question_rows(book["Ques 1"])
    questions2 = question_rows(book["Ques 2"])
    questions3 = question_rows(book["Ques 3"])
    questions4 = question_rows(book["Ques 4"])
    validate_question_graph([questions1, questions2, questions3, questions4])

    heads = []
    head_sheet = book['diagram head']
    head_columns = required_columns(head_sheet, ['electrical diagram 1', 'Ключ на схеме Э3', 'key_diagram'])
    known_keys = {node['diagramKey'] for node in questions1 + questions2 + questions3 + questions4 if node['diagramKey']}
    for source_row, row in enumerate(rows(head_sheet)[1:], start=2):
        code = cell(row, head_columns, 'electrical diagram 1')
        if not code:
            continue
        placeholder = cell(row, head_columns, 'Ключ на схеме Э3')
        key = cell(row, head_columns, 'key_diagram')
        if not isinstance(placeholder, str) or not placeholder.startswith('#') or key not in known_keys:
            raise ValueError(f'diagram head:{source_row}: неверный ключ или поле схемы')
        if any(item['code'] == code and item['keyDiagram'] == key for item in heads):
            raise ValueError(f'diagram head:{source_row}: повтор ключа {key} для {code}')
        heads.append({'sourceRow': source_row, 'code': code, 'placeholder': placeholder, 'keyDiagram': key})

    index1 = []
    for excel_row, row in enumerate(rows(book["Load index 1"])[1:], start=2):
        if row[8] is None:
            continue
        index1.append({
            "sourceRow": excel_row, "voltageCode": value(row[0]), "voltage": value(row[1]),
            "startCode": value(row[2]), "start": value(row[3]), "tripCode": value(row[4]),
            "trip": value(row[5]), "currentFrom": num(row[6]), "currentTo": num(row[7]),
            "loadIndex": value(row[8]), "linesAmount": value(row[9]),
        })

    index2 = []
    for excel_row, row in enumerate(rows(book["Load index 2"])[1:], start=2):
        if row[8] is None:
            continue
        index2.append({
            "sourceRow": excel_row, "selectionCode": value(row[0]), "selection": value(row[1]),
            "voltageCode": value(row[2]), "voltage": value(row[3]), "startCode": value(row[4]),
            "start": value(row[5]), "currentFrom": num(row[6]), "currentTo": num(row[7]),
            "loadIndex": value(row[8]), "linesAmount": value(row[9]),
        })

    diagrams1 = []
    d1 = book["diagram 1"]
    c1 = required_columns(d1, ["Motor starting / control method", "Location (VFD start / Triac speed control)",
        "Rated voltage, V", "Circuit breaker trip type", "controller_Manufacturer (Brand)",
        "controller_Type", "electrical diagram 1", "electrical diagram 2", "lines_amount",
        "consumption_current", *IO_CODES])
    for excel_row, row in enumerate(rows(d1)[1:], start=2):
        if cell(row, c1, "electrical diagram 1") is None:
            continue
        diagrams1.append({
            "sourceRow": excel_row, "startCode": cell(row, c1, "Motor starting / control method"),
            "start": cell(row, c1, "Способ пуска / управления двигателем"),
            "locationCode": cell(row, c1, "Location (VFD start / Triac speed control)"),
            "location": cell(row, c1, "Расположение (ЧП / Симисторный регулятор скорости)"),
            "voltageCode": cell(row, c1, "Rated voltage, V"), "voltage": cell(row, c1, "Номинальное напряжение, В"),
            "tripCode": cell(row, c1, "Circuit breaker trip type"), "trip": cell(row, c1, "Тип защиты автоматического выключателя"),
            "controller": cell(row, c1, "controller_Manufacturer (Brand)"), "controllerType": cell(row, c1, "controller_Type"),
            "diagram1": cell(row, c1, "electrical diagram 1"), "diagram2": cell(row, c1, "electrical diagram 2"),
            "io": io(row, c1, excel_row), "linesAmount": cell(row, c1, "lines_amount"),
            "consumptionCurrent": cell(row, c1, "consumption_current"), "placement": cell(row, c1, "diagram placement"),
        })

    diagrams2 = []
    d2 = book["diagram 2"]
    c2 = required_columns(d2, ["Thermal overload contact (bimetallic)", "PTC thermistor protection",
        "Winding RTD (Pt100)", "Connection type_Bearing temperature sensor (RTD) (Pt100)",
        "controller_Manufacturer (Brand)", "controller_Type", "electrical diagram 1",
        "electrical diagram 2", "lines_amount", "consumption_current", *IO_CODES])
    for excel_row, row in enumerate(rows(d2)[1:], start=2):
        if cell(row, c2, "electrical diagram 1") is None:
            continue
        diagrams2.append({
            "sourceRow": excel_row, "thermalCode": cell(row, c2, "Thermal overload contact (bimetallic)"),
            "thermal": value(row[c2["Thermal overload contact (bimetallic)"] + 1]),
            "ptcCode": cell(row, c2, "PTC thermistor protection"), "ptc": value(row[c2["PTC thermistor protection"] + 1]),
            "windingWireCode": cell(row, c2, "Winding RTD (Pt100)"), "windingWire": value(row[c2["Winding RTD (Pt100)"] + 1]),
            "bearingWireCode": cell(row, c2, "Connection type_Bearing temperature sensor (RTD) (Pt100)"),
            "bearingWire": value(row[c2["Connection type_Bearing temperature sensor (RTD) (Pt100)"] + 1]),
            "controller": cell(row, c2, "controller_Manufacturer (Brand)"), "controllerType": cell(row, c2, "controller_Type"),
            "diagram1": cell(row, c2, "electrical diagram 1"), "diagram2": cell(row, c2, "electrical diagram 2"),
            "io": io(row, c2, excel_row), "linesAmount": cell(row, c2, "lines_amount"),
            "consumptionCurrent": cell(row, c2, "consumption_current"), "placement": cell(row, c2, "diagram placement"),
        })

    manualVfd = []
    for excel_row, row in enumerate(rows(book["sp1"])[1:], start=2):
        if row[4] is None:
            continue
        manualVfd.append({"sourceRow": excel_row, "location": value(row[1]), "selectionCode": value(row[2]), "selection": value(row[3]), "model": value(row[4]), "loadIndex": value(row[5]), "name": value(row[6]), "quantity": value(row[7]), "unit": value(row[8]), "lde": value(row[9])})

    automaticVfd = []
    for excel_row, row in enumerate(rows(book["sp2"])[1:], start=2):
        if row[4] is None:
            continue
        automaticVfd.append({"sourceRow": excel_row, "location": value(row[1]), "selectionCode": value(row[2]), "selection": value(row[3]), "brand": value(row[4]), "loadIndex": value(row[5]), "name": value(row[6]), "quantity": value(row[7]), "unit": value(row[8]), "lde": value(row[9])})

    sensorSpecification = []
    for excel_row, row in enumerate(rows(book["sp3"])[1:], start=2):
        if row[0] is None:
            continue
        sensorSpecification.append({"sourceRow": excel_row, "diagram1": value(row[0]), "name": value(row[1]), "quantity": value(row[2]), "unit": value(row[3]), "lde": value(row[4])})

    mainSpecification = []
    for excel_row, row in enumerate(rows(book["sp4"])[1:], start=2):
        if row[0] is None:
            continue
        mainSpecification.append({"sourceRow": excel_row, "diagram1": value(row[0]), "loadIndex": value(row[1]), "name": value(row[2]), "quantity": value(row[3]), "unit": value(row[4]), "lde": value(row[5])})

    target = args.target.expanduser().resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps({
        "rulesSchemaVersion": 3,
        "source": {"file": source.name, "revision": args.revision or source.stem, "sha256": hashlib.sha256(source.read_bytes()).hexdigest(), "sheets": ["Ques 1", "Ques 2", "Ques 3", "Ques 4", "diagram head", "Load index 1", "Load index 2", "diagram 1", "diagram 2", "sp1", "sp2", "sp3", "sp4"]},
        "questions1": questions1, "questions2": questions2, "questions3": questions3, "questions4": questions4, "loadIndex1": index1, "loadIndex2": index2, "diagrams1": diagrams1, "diagrams2": diagrams2,
        "diagramHeads": heads,
        "manualVfd": manualVfd, "automaticVfd": automaticVfd, "sensorSpecification": sensorSpecification, "mainSpecification": mainSpecification,
    }, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"Wrote {target} | questions {len(questions1)}/{len(questions2)}/{len(questions3)}/{len(questions4)} | diagrams {len(diagrams1)}/{len(diagrams2)}")


if __name__ == "__main__":
    main()
