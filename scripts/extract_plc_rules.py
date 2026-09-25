#!/usr/bin/env python3
"""Build a versioned PLC pin catalog from the supplied engineering workbooks."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCES = ROOT / 'data' / 'rules-source'
COMPOSITION = SOURCES / 'Состав ПЛК.xlsx'
MARKING = SOURCES / 'Правило_обозначения и маркировки элемкнтов на схеме_Э3_В2.xlsx'
TARGET = ROOT / 'data' / 'plc-catalog.json'


def records(sheet):
    rows = list(sheet.iter_rows(values_only=True))
    columns = {str(name).strip(): index for index, name in enumerate(rows[0]) if name is not None}
    return columns, [tuple(row) + (None,) * (len(rows[0]) - len(row)) for row in rows[1:]]


def main():
    composition = load_workbook(COMPOSITION, read_only=True, data_only=True)
    marking = load_workbook(MARKING, read_only=True, data_only=True)
    base_columns, base_rows = records(composition['base_controller'])
    required = {'controller_Manufacturer (Brand)', 'controller_Type', 'general signal amount', 'max_moduls'}
    if not required <= set(base_columns) or len(base_rows) != 1:
        raise ValueError('Ожидается одна описанная базовая модель ПЛК')
    base = base_rows[0]
    brand = base[base_columns['controller_Manufacturer (Brand)']]
    model = base[base_columns['controller_Type']]
    code_columns = [name for name in base_columns if name not in required | {'consumption_current'}]
    limits = {code: int(base[base_columns[code]]) for code in code_columns}
    needed = {'controller_Manufacturer (Brand)', 'controller_Type', 'signal code', 'CND_COM', 'entrance number 1',
              'entrance number 2', 'entrance number 3', 'entrance number 4', 'group',
              'Ключ на схеме Э3', 'Ключ для общего GND на схеме Э3'}
    def read_options(sheet_name, expected_model, expected_limits):
        key_columns, key_rows = records(marking[sheet_name])
        if not needed <= set(key_columns):
            raise ValueError(f'{sheet_name}: не хватает колонок ключей зажимов')
        options = []
        for source_row, row in enumerate(key_rows, start=2):
            if row[key_columns['controller_Manufacturer (Brand)']] != brand or row[key_columns['controller_Type']] != expected_model:
                raise ValueError(f'{sheet_name}:{source_row}: другая модель или изготовитель')
            code = row[key_columns['signal code']]
            if code not in expected_limits:
                raise ValueError(f'{sheet_name}:{source_row}: неизвестный тип сигнала {code}')
            terminals = [str(row[key_columns[f'entrance number {index}']]) for index in range(1, 5)
                         if row[key_columns[f'entrance number {index}']] is not None]
            if not terminals:
                raise ValueError(f'{sheet_name}:{source_row}: нет номера вывода')
            key = row[key_columns['Ключ на схеме Э3']]
            if key != '#' + code:
                raise ValueError(f'{sheet_name}:{source_row}: ключ {key} не соответствует {code}')
            common = row[key_columns['CND_COM']]
            common_key = row[key_columns['Ключ для общего GND на схеме Э3']]
            if common_key != ('#' + common if common else None):
                raise ValueError(f'{sheet_name}:{source_row}: CND_COM не совпадает с ключом общего вывода')
            options.append({'sourceRow': source_row, 'code': code, 'terminals': terminals,
                            'group': row[key_columns['group']], 'key': key,
                            'commonKey': common_key, 'commonDesignation': common})
        for code, limit in expected_limits.items():
            count = len({item['terminals'][0] for item in options if item['code'] == code})
            if count != limit:
                raise ValueError(f'{sheet_name}: {code}: состав {limit}, ключи {count}')
        if len({(item['code'], item['terminals'][0]) for item in options}) != len(options):
            raise ValueError(f'{sheet_name}: повтор тип сигнала / вывод')
        return options

    options = read_options('Ключи зажимов ПЛК', model, limits)
    module_columns, module_rows = records(composition['moduls'])
    if len(module_rows) != 1 or not {'controller_Manufacturer (Brand)', 'controller_Type', 'general signal amount', *code_columns} <= set(module_columns):
        raise ValueError('Ожидается один описанный тестовый модуль')
    module_row = module_rows[0]
    module_model = module_row[module_columns['controller_Type']]
    if module_row[module_columns['controller_Manufacturer (Brand)']] != brand:
        raise ValueError('Изготовитель модуля не совпал с контроллером')
    module_limits = {code: int(module_row[module_columns[code]]) for code in code_columns}
    module_options = read_options('Клеммы модуля тест', module_model, module_limits)
    template_columns, template_rows = records(marking['Поля схем Э3'])
    template_required = {'electrical diagram 1', 'role', 'Ключ в DXF', 'signal code', 'Значение из строки вывода'}
    if not template_required <= set(template_columns):
        raise ValueError('Поля схем Э3: не хватает колонок привязки CAD')
    template_fields = {}
    seen_roles = set()
    for source_row, row in enumerate(template_rows, start=2):
        code = row[template_columns['electrical diagram 1']]
        role = row[template_columns['role']]
        marker = row[template_columns['Ключ в DXF']]
        signal = row[template_columns['signal code']]
        value_from = row[template_columns['Значение из строки вывода']]
        if (not isinstance(code, str) or not code.startswith('im1-')
                or role not in {'plc.di', 'plc.common', 'plc.do.1', 'plc.do.2', 'plc.do.common'}
                or signal not in limits
                or value_from not in {'entrance number 1', 'entrance number 2', 'CND_COM'}
                or not isinstance(marker, str) or not marker.startswith('#')
                or (code, role) in seen_roles):
            raise ValueError(f'Поля схем Э3:{source_row}: неверная или повторная привязка')
        if (role.endswith('common') and value_from != 'CND_COM') or (not role.endswith('common') and value_from == 'CND_COM'):
            raise ValueError(f'Поля схем Э3:{source_row}: роль не соответствует источнику значения')
        if not role.endswith('common') and marker != '#' + signal:
            raise ValueError(f'Поля схем Э3:{source_row}: ключ сигнала не соответствует {signal}')
        seen_roles.add((code, role))
        template_fields.setdefault(code, []).append({'sourceSheet': 'Поля схем Э3', 'sourceRow': source_row,
                                                       'role': role, 'marker': marker, 'signal': signal,
                                                       'valueFrom': value_from})
    if int(base[base_columns['max_moduls']]) < 0:
        raise ValueError('Некорректное максимальное число модулей')
    group_columns, group_rows = records(composition['add_cond'])
    group_required = {'controller_Manufacturer (Brand)', 'controller_Type', 'group', 'signal amount', *code_columns}
    if not group_required <= set(group_columns):
        raise ValueError('Не хватает колонок в add_cond')
    for catalog_model, catalog_options, amount in ((model, options, base[base_columns['general signal amount']]),
                                                    (module_model, module_options, module_row[module_columns['general signal amount']])):
        seen_groups = set()
        for source_row, row in enumerate(group_rows, start=2):
            if row[group_columns['controller_Manufacturer (Brand)']] != brand or row[group_columns['controller_Type']] != catalog_model:
                continue
            group = row[group_columns['group']]
            if group in seen_groups:
                raise ValueError(f'add_cond:{source_row}: повтор группы {group}')
            seen_groups.add(group)
            listed = [item for item in catalog_options if item['group'] == group]
            count = len({item['terminals'][0] for item in listed})
            if count != row[group_columns['signal amount']]:
                raise ValueError(f'add_cond:{source_row}: группа {group} содержит {count} выводов вместо {row[group_columns["signal amount"]]}')
            for code in code_columns:
                actual = len([item for item in listed if item['code'] == code])
                expected = row[group_columns[code]] or 0
                if actual != expected:
                    raise ValueError(f'add_cond:{source_row}: {code} в группе {group}: {actual} вместо {expected}')
        if {item['group'] for item in catalog_options} != seen_groups:
            raise ValueError(f'{catalog_model}: ключи содержат группу вне add_cond')
        if len({item['terminals'][0] for item in catalog_options}) != amount:
            raise ValueError(f'{catalog_model}: число уникальных ресурсов не совпало с составом')
    payload = {
        'schemaVersion': 2,
        'sources': {
            'composition': {'file': COMPOSITION.name, 'sha256': hashlib.sha256(COMPOSITION.read_bytes()).hexdigest()},
            'marking': {'file': MARKING.name, 'sha256': hashlib.sha256(MARKING.read_bytes()).hexdigest()},
        },
        'controller': {'brand': brand, 'model': model, 'generalSignalAmount': base[base_columns['general signal amount']],
                       'maxModules': base[base_columns['max_moduls']], 'limits': limits},
        'pinOptions': options,
        'module': {'brand': brand, 'model': module_model,
                   'generalSignalAmount': module_row[module_columns['general signal amount']],
                   'limits': module_limits, 'terminalStatus': 'testDerived'},
        'modulePinOptions': module_options,
        'templateFields': template_fields,
    }
    TARGET.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f'Wrote {TARGET} | {brand} {model} | {len(options)} pin options')


if __name__ == '__main__':
    main()
